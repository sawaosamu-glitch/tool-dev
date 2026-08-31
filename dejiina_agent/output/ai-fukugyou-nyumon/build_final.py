#!/usr/bin/env python3
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

base_dir = Path("output/ai-fukugyou-nyumon")
md_path = base_dir / "manuscript.md"
out_path = base_dir / "final_book.docx"

NORMAL_MARGIN_CM = 2.0
RED_BOLD = RGBColor(0xCC, 0x22, 0x00)
BLUE_H1 = RGBColor(0x4A, 0x6C, 0xF7)
CONTENT_IMG_WIDTH = (14.8 - NORMAL_MARGIN_CM * 2) / 2.54

NUMBER_PATTERN = re.compile(
    r'\d+(?:[,，]\d+)*(?:\.\d+)?[%％円万千倍冊日件本時間分個人]+'
    r'|\d+(?:[,，]\d+)+'
)

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(14.8)
sec.page_height = Cm(21.0)
sec.left_margin = sec.right_margin = Cm(NORMAL_MARGIN_CM)
sec.top_margin = sec.bottom_margin = Cm(NORMAL_MARGIN_CM)

doc.styles['Normal'].font.name = 'Meiryo'
doc.styles['Normal'].font.size = Pt(10.5)

img_pattern = re.compile(r'!\[([^\]]*)\]\(images/([^)]+)\)(\{[^}]*\})?')

def add_heading(text, level):
    p = doc.add_heading(text, level=level)
    if p.runs:
        p.runs[0].font.name = 'Meiryo'
        colors = {1: BLUE_H1, 2: RGBColor(0x33,0x33,0x33), 3: RGBColor(0x55,0x55,0x55)}
        sizes = {1: Pt(16), 2: Pt(13), 3: Pt(11)}
        p.runs[0].font.color.rgb = colors.get(level, RGBColor(0,0,0))
        p.runs[0].font.size = sizes.get(level, Pt(11))
    return p

def add_text_paragraph(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    bold_parts = re.split(r'\*\*(.+?)\*\*', text)
    for j, part in enumerate(bold_parts):
        if not part:
            continue
        if j % 2 == 1:
            run = p.add_run(part)
            run.font.name = 'Meiryo'
            run.font.size = Pt(10.5)
            run.bold = True
            run.font.color.rgb = RED_BOLD
        else:
            sub_parts = NUMBER_PATTERN.split(part)
            num_matches = NUMBER_PATTERN.findall(part)
            for k, sub in enumerate(sub_parts):
                if sub:
                    run = p.add_run(sub)
                    run.font.name = 'Meiryo'
                    run.font.size = Pt(10.5)
                if k < len(num_matches):
                    run = p.add_run(num_matches[k])
                    run.font.name = 'Meiryo'
                    run.font.size = Pt(10.5)
                    run.bold = True
                    run.font.color.rgb = RED_BOLD
    return p

content = md_path.read_text(encoding="utf-8")
lines = content.split("\n")

first_para = True
for line in lines:
    if line.strip() == '\\newpage':
        doc.add_page_break()
        continue

    img_m = img_pattern.match(line.strip())
    if img_m:
        img_path = base_dir / "images" / img_m.group(2)
        if img_path.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.add_run().add_picture(str(img_path), width=Inches(CONTENT_IMG_WIDTH))
        continue

    if line.startswith('# '):
        if not first_para:
            pass
        add_heading(line[2:].strip(), 1)
    elif line.startswith('## '):
        add_heading(line[3:].strip(), 2)
    elif line.startswith('### '):
        add_heading(line[4:].strip(), 3)
    elif line.strip().startswith('```'):
        continue
    elif line.strip() not in ('---', ''):
        text = line.strip()
        if text:
            sentences = re.split(r'(?<=。)', text)
            sentences = [s for s in sentences if s.strip()]
            if len(sentences) > 1:
                for sentence in sentences:
                    if sentence.strip():
                        add_text_paragraph(sentence.strip())
            else:
                add_text_paragraph(text)

    first_para = False

doc.save(str(out_path))
size_kb = out_path.stat().st_size // 1024
print(f"[OK] final_book.docx ({size_kb} KB) → {out_path}")
